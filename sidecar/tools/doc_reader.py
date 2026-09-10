"""文档解析：把 .docx/.xlsx/.pptx/.pdf 提取为可读文本 + 格式概要。

## 为什么需要这个模块（0.4.19 #4）

read_file 此前对所有非图片文件一律执行
`read_bytes().decode("utf-8", errors="replace")`。
.docx/.pptx/.xlsx 是 **zip 压缩包**、.pdf 是**二进制格式**，这样解出来必然是
`PK\x03\x04...` 一类的乱码（用户实测：Agent 只能看到二进制垃圾，完全读不出
上传的参考文档）。图片早有 base64 特例（checkpoint-067 R-4）让多模态模型真正
"看图"，Office/PDF 却没有对应通道——**这是应用的能力缺口，不是模型不行**。

本模块补上这条通道：解析出**内容文本（含表格）**与**格式概要**，
使 Agent 能真正读取用户给的参考文档，用于"按原格式产出"这类任务（#14）。

## 设计取舍

1. **表格必须提取**。只遍历 `document.paragraphs` 会丢掉全部表格；而用户参考文档
   （律师函、证据清单）里表格恰恰是格式关键。故 docx 按 body 的 XML 子元素顺序
   交替取段落与表格，保持原始排版次序。
2. **格式信息聚合而非逐段罗列**（#14）。逐段输出字体/字号会让上下文瞬间爆炸
   （几百段就是几万 token）。改为按「样式 × 字体 × 字号 × 加粗 × 对齐 × 首行缩进」
   去重聚合，只列实际出现的组合及其段数——Agent 据此足以判断
   "标题=黑体16pt居中、正文=宋体12pt首行缩进0.74cm"，复刻格式够用。
3. **中文字体要读 eastAsia**。`python-docx` 的 `font.name` 只取 ascii 字体
   （`w:rFonts/@w:ascii`），中文文档的宋体/黑体挂在 `@w:eastAsia` 上，
   不另取就恒为 None——对中文参考文档等于没提取到格式。
4. **全部延迟 import**（与 doc_writer 一致）。read_file 是通用工具，
   不能因为缺一个解析库就整体不可用；缺库时返回明确提示。
5. **pptx 走 zip 兜底**。项目未安装 python-pptx，但 pptx 与 docx 同为 OOXML
   zip 包、文本都在 `<a:t>` 节点里 → 直接解 zip 提取，不为此引入新依赖。

## 依赖现状（sidecar/.venv，2026-09-10 实测）

| 类型 | 库 | 状态 |
|---|---|---|
| docx | python-docx 1.2.0 | ✅ 已装 |
| xlsx | openpyxl 3.1.5 | ✅ 已装 |
| pdf | pypdf 6.16.2 | ✅ 已装 |
| pptx | python-pptx | ❌ 未装 → zip 兜底 |
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any

# 可由本模块解析的扩展名（registry.read_file 据此决定走解析分支还是原字节解码）
DOC_EXTS = {".docx", ".xlsx", ".xlsm", ".pptx", ".pdf"}

# 旧式 .doc/.xls/.ppt 是 OLE 复合二进制，python-docx/openpyxl/pypdf 都读不了。
# ⛔ 但 .doc 在 macOS 上**不是死路**：系统自带 /usr/bin/textutil 能把它转成 .docx，
# 再交 python-docx 解析——零外部依赖、零体积增量（不要引入 antiword/libreoffice）。
# textutil 不支持 .xls/.ppt，那两个才真的只能提示另存。
LEGACY_EXTS = {".doc", ".xls", ".ppt"}
# textutil 支持的旧格式 → 目标新格式（其余扩展名它不认，不要往这里加）
_TEXTUTIL_TARGET = {".doc": "docx", ".rtf": "docx", ".odt": "docx"}

_PPTX_TEXT_RE = re.compile(r"<a:t>(.*?)</a:t>", re.S)
# docx 段落级 eastAsia 字体：w:rFonts/@w:eastAsia
_EASTASIA_RE = re.compile(r'w:eastAsia="([^"]+)"')

# WD_ALIGN_PARAGRAPH 枚举值 → 人话
_ALIGN_LABEL = {0: "左对齐", 1: "居中", 2: "右对齐", 3: "两端对齐", 4: "分散对齐"}


def is_parseable(path: Path) -> bool:
    """该扩展名是否走文档解析分支（含旧式 Office，内部再按平台判断能否处理）。"""
    ext = path.suffix.lower()
    return ext in DOC_EXTS or ext in LEGACY_EXTS


def is_legacy_office(path: Path) -> bool:
    return path.suffix.lower() in LEGACY_EXTS


def _convert_legacy(path: Path) -> tuple[Any, Path, str] | None:
    """旧式 Office → OOXML 临时文件，复用既有解析链。

    返回 (临时目录对象, 转换后路径, 小写扩展名)；不可行时返回 None
    （平台无 textutil、格式不在其支持列表、或转换命令失败）。

    ⛔ 用 macOS 自带 textutil，**不要**引入 antiword / textract / libreoffice：
    零外部依赖、零体积增量，且与 PyInstaller 打包无冲突。
    ⛔ 只读源文件、只写临时目录——绝不修改或覆盖用户的原始文件。
    ⛔ 临时目录由调用方 cleanup()（TemporaryDirectory 上下文对象）。
    """
    ext = path.suffix.lower()
    target_fmt = _TEXTUTIL_TARGET.get(ext)
    if not target_fmt:
        return None
    if not shutil.which("textutil"):
        return None  # Windows/Linux 无此命令 → 由调用方给出"另存为新格式"提示
    try:
        tmp = tempfile.TemporaryDirectory(prefix="docreader_")
        out = Path(tmp.name) / (path.stem + "." + target_fmt)
        # timeout 防坏文件导致命令挂死（read_file 是同步工具，挂住会拖住整轮）
        proc = subprocess.run(
            ["textutil", "-convert", target_fmt, str(path), "-output", str(out)],
            capture_output=True, timeout=60,
        )
        if proc.returncode != 0 or not out.exists() or out.stat().st_size == 0:
            tmp.cleanup()
            return None
        return tmp, out, "." + target_fmt
    except Exception:
        return None


def extract(path: Path, budget: int) -> dict[str, Any]:
    """解析文档，返回可直接并入 read_file 结果的字段。

    budget：content 的 UTF-8 字节预算（调用方已扣掉头部提示语的开销）。
    返回 {"ok", "content", "parse_kind", "parse_error"}：
      ok=True  → content 为解析出的文本（格式概要在前、正文在后）
      ok=False → content 为**说明性文字**（不是乱码），parse_error 为机器可读原因

    ⛔ 失败也返回说明文字而非抛异常：read_file 的调用方是模型，
    给它一句"这个格式读不了，因为 X，建议 Y"远比给它一个 500 有用。
    """
    orig = path  # ⛔ 提示文案一律用原文件名（转换后的临时名对用户毫无意义）
    ext = path.suffix.lower()
    detail = ""
    body = ""
    _tmpdir: Any = None
    try:
        if ext in LEGACY_EXTS:
            _conv = _convert_legacy(path)
            if _conv is None:
                # textutil 不支持该格式（.xls/.ppt）或平台无此命令 → 如实提示，绝不编造
                _new = {"doc": "docx", "xls": "xlsx", "ppt": "pptx"}.get(ext.lstrip("."), "新格式")
                _plat = "" if shutil.which("textutil") else "（当前平台无 textutil，该能力仅 macOS 可用）"
                return {"ok": False,
                        "content": (f"文件 {orig.name} 是旧式 {ext} 格式，当前环境无法直接解析{_plat}。"
                                    f"请如实告知用户：需要另存为 .{_new} 后才能读取，"
                                    f"不要编造文件内容。"),
                        "parse_kind": ext.lstrip("."), "parse_error": f"unsupported_legacy:{ext}"}
            _tmpdir, path, ext = _conv  # 转换产物按其真实类型继续走下面的解析链
        if ext == ".docx":
            body, detail = _read_docx(path)
        elif ext in (".xlsx", ".xlsm"):
            body, detail = _read_xlsx(path)
        elif ext == ".pptx":
            body, detail = _read_pptx(path)
        elif ext == ".pdf":
            body, detail = _read_pdf(path)
        else:
            return {"ok": False, "content": f"暂不支持解析 {ext} 格式。",
                    "parse_kind": ext.lstrip("."), "parse_error": f"unsupported_ext:{ext}"}
    except ImportError as e:
        return {"ok": False,
                "content": (f"解析 {orig.name} 需要 Python 库，但当前环境未安装（{e}）。"
                            f"请如实告知用户无法读取该文件，不要编造内容。"),
                "parse_kind": ext.lstrip("."), "parse_error": f"missing_dependency:{e}"}
    except zipfile.BadZipFile:
        return {"ok": False,
                "content": (f"文件 {orig.name} 不是有效的 {ext} 文档（zip 结构损坏，"
                            f"或扩展名与实际格式不符）。请如实告知用户，不要编造内容。"),
                "parse_kind": ext.lstrip("."), "parse_error": "corrupt_or_not_ooxml"}
    except Exception as e:  # 解析失败不能拖垮 read_file
        return {"ok": False,
                "content": (f"解析 {orig.name} 时出错（{type(e).__name__}: {e}）。"
                            f"请如实告知用户无法读取，不要编造内容。"),
                "parse_kind": ext.lstrip("."), "parse_error": f"parse_error:{type(e).__name__}"}
    finally:
        # 旧格式转换的临时目录必须清理（textutil 产物是中间件，不该留在磁盘）
        if _tmpdir is not None:
            try:
                _tmpdir.cleanup()
            except Exception:
                pass

    text = body.strip()
    # 组装：格式概要在前（截断时优先保住它——复刻格式靠的是概要，不是正文尾部）
    full = (detail.strip() + "\n\n" + text) if detail.strip() else text
    if not full.strip():
        return {"ok": True,
                "content": f"文件 {orig.name} 已解析，但未提取到任何文本（可能是空文档，"
                           f"或内容为图片/扫描型，需走图像识别）。",
                "parse_kind": ext.lstrip("."), "parse_error": ""}

    raw = full.encode("utf-8")
    truncated = len(raw) > budget
    if truncated:
        full = raw[:budget].decode("utf-8", errors="ignore").rstrip() + "\n\n…（内容过长，已截断）"
    return {"ok": True, "content": full, "parse_kind": ext.lstrip("."),
            "parse_error": "", "truncated": truncated}


# ────────────────────────────── docx ──────────────────────────────

def _iter_docx_blocks(doc: Any):
    """按文档流顺序交替产出 Paragraph / Table。

    ⛔ 不能用 doc.paragraphs + doc.tables 分别遍历再拼接——那样表格会全部
    堆到正文末尾，丢失原始排版次序（用户参考文档里"表格夹在段落之间"
    是常见版式，顺序错了 Agent 复刻出来的格式就是错的）。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _docx_run_font(run: Any) -> tuple[str | None, float | None, bool | None]:
    """取一个 run 的字体名/字号/加粗。字体名须兼顾 ascii 与 eastAsia（中文）。"""
    name: str | None = None
    size: float | None = None
    bold: bool | None = None
    try:
        f = run.font
        name = f.name
        if f.size is not None:
            size = round(f.size.pt, 1)
        if f.bold is not None:
            bold = bool(f.bold)
        # python-docx 的 font.name 只映射 w:ascii；中文字体在 w:eastAsia
        rpr = getattr(run._element, "rPr", None)
        if rpr is not None:
            ea = rpr.get("eastAsia") if hasattr(rpr, "get") else None
            if not ea:
                m = _EASTASIA_RE.search(rpr.xml if hasattr(rpr, "xml") else "")
                ea = m.group(1) if m else None
            if ea:
                name = ea if (not name or name == ea) else f"{name}/{ea}"
    except Exception:
        pass
    return name, size, bold


def _align_label(al: Any) -> str | None:
    if al is None:
        return None
    try:
        return _ALIGN_LABEL.get(int(al), str(al))
    except Exception:
        return str(al)


def _docx_image_count(path: Path) -> tuple[int, int]:
    """统计 docx 内嵌图片数与字节数。

    ⛔ 必须有这个：实测 7.6MB 的「证据（王永斌）.docx」是图片为主的文档，
    文本只有几十字。若只回文本，模型会误判"这文件基本是空的"，
    进而凭空编造或反复重试——如实告知"含 N 张图、共 X MB，需走图像识别"才有用。
    """
    total = 0
    nbytes = 0
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith("word/media/"):
                    try:
                        total += 1
                        nbytes += z.getinfo(n).file_size
                    except Exception:
                        continue
    except Exception:
        return 0, 0
    return total, nbytes


def _read_docx(path: Path) -> tuple[str, str]:
    import docx

    d = docx.Document(str(path))
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for block in _iter_docx_blocks(d):
        if isinstance(block, Paragraph):
            text = (block.text or "").strip()
            if not text:
                continue
            style = ""
            try:
                style = (block.style.name if block.style is not None else "") or ""
            except Exception:
                style = ""
            # 标题按层级转 markdown #，让 Agent 看清文档结构骨架
            low = style.lower()
            if low.startswith("heading") or style.startswith("标题"):
                m = re.search(r"(\d+)", style)
                lvl = min(int(m.group(1)) if m else 1, 6)
                lines.append("#" * lvl + " " + text)
            else:
                lines.append(text)
        elif isinstance(block, Table):
            rows: list[str] = []
            for row in block.rows:
                cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                # 首行按表头处理，补 markdown 分隔行，渲染/阅读都清晰
                width = rows[0].count("|") - 1
                lines.append("")
                lines.append(rows[0])
                lines.append("|" + "---|" * max(width, 1))
                lines.extend(rows[1:])
                lines.append("")

    return "\n".join(lines), _docx_format_summary(d, path)


def _docx_format_summary(d: Any, path: Path | None = None) -> str:
    """文档级 + 样式级格式概要（聚合去重，防上下文爆炸）。"""
    out: list[str] = ["【格式概要】"]

    # ── 内嵌图片 ──（含图为主的文档必须提示，否则模型会误判文件是空的）
    if path is not None:
        _nimg, _nbytes = _docx_image_count(path)
        if _nimg:
            out.append(f"内嵌图片：{_nimg} 张（约 {round(_nbytes / 1048576, 2)}MB）。"
                       f"⛔ 图片内容无法由文本解析获得，若任务需要图中文字/信息，"
                       f"必须走图像识别（read_file 逐张读图或委派视觉子 Agent），"
                       f"不得据本文本臆测图片内容。")

    # ── 页面设置 ──
    try:
        s = d.sections[0]

        def _cm(v: Any) -> Any:
            return round(v.cm, 2) if v is not None else None

        pw, ph = _cm(s.page_width), _cm(s.page_height)
        orient = "横向" if (pw and ph and pw > ph) else "纵向"
        out.append(f"页面：{orient}，{pw}×{ph}cm；页边距 上{_cm(s.top_margin)} "
                   f"下{_cm(s.bottom_margin)} 左{_cm(s.left_margin)} 右{_cm(s.right_margin)}cm")
    except Exception:
        pass

    # ── 段落格式聚合 ──
    combos: dict[tuple, int] = {}
    for p in d.paragraphs:
        if not (p.text or "").strip():
            continue
        try:
            style = (p.style.name if p.style is not None else "") or "Normal"
        except Exception:
            style = "Normal"
        name, size, bold = _docx_run_font(p.runs[0]) if p.runs else (None, None, None)
        fli = None
        lsp = None
        try:
            pf = p.paragraph_format
            if pf.first_line_indent is not None:
                fli = round(pf.first_line_indent.cm, 2)
            # ⛔ line_spacing 有两种语义，必须区分（实测踩坑）：
            #   · float（如 1.5）= **倍数**行距 → 原样输出
            #   · Length（有 .pt）= **固定**行距，其 int 值是 EMU（22磅 → 279400）
            #     直接输出就是"行距=279400.0"这种鬼数字，模型无法据此复刻排版。
            _ls = pf.line_spacing
            if _ls is not None:
                lsp = (f"{round(float(_ls.pt), 1)}磅" if hasattr(_ls, "pt")
                       else f"{round(float(_ls), 2)}倍")
        except Exception:
            pass
        key = (style, name, size, bold, _align_label(p.alignment), fli, lsp)
        combos[key] = combos.get(key, 0) + 1

    if combos:
        out.append("段落格式（样式 / 字体 / 字号pt / 加粗 / 对齐 / 首行缩进cm / 行距 → 段数）：")
        # 按出现次数降序，最多 12 行——覆盖主格式即可，长尾不值得占上下文
        for (style, name, size, bold, al, fli, lsp), cnt in sorted(
                combos.items(), key=lambda kv: -kv[1])[:12]:
            parts = [style or "Normal",
                     f"字体={name}" if name else "字体=继承",
                     f"字号={size}" if size else "字号=继承",
                     f"加粗={'是' if bold else '否'}" if bold is not None else "",
                     f"对齐={al}" if al else "",
                     f"首行缩进={fli}" if fli else "",
                     f"行距={lsp}" if lsp else ""]
            out.append("  " + " / ".join(x for x in parts if x) + f" → {cnt} 段")

    # ── 规模统计 ──
    try:
        ntbl = len(d.tables)
        if ntbl:
            out.append(f"表格：{ntbl} 个")
    except Exception:
        pass

    return "\n".join(out) if len(out) > 1 else ""


# 对齐 label（人话）→ WD_ALIGN_PARAGRAPH 枚举值（写出端套用参考格式时反查）
# ⛔ 与上方 _ALIGN_LABEL 互为逆映射；新增对齐方式时两处必须同步。
_ALIGN_VALUE = {v: k for k, v in _ALIGN_LABEL.items()}


def extract_docx_style(path: Path) -> dict[str, Any]:
    """提取参考 .docx 的**结构化**格式，供 doc_writer 套用到新生成的文档（#14 写出端）。

    与 `_docx_format_summary` 的区别：后者返回**给人/模型看的字符串概要**（聚合去重、
    带中文说明），本函数返回**机器可读的 dict**，字段是 doc_writer 能直接消费的原值。

    返回（任一项缺失即为 None / 空 dict，写出端据此回退到自身默认值）：
      {"page": {"width_cm","height_cm","top_cm","bottom_cm","left_cm","right_cm"},
       "body": {"font","ascii_font","size_pt","align_value","first_line_indent_cm",
                "line_spacing"(float 倍数 或 None),"line_spacing_pt"(float 磅 或 None)}}

    ⛔ body 取「正文主格式」：遍历段落，按出现次数最多的 (字体,字号,对齐,缩进,行距) 组合
      判定为正文格式——标题/页脚的少数格式不应主导正文排版。与 _docx_format_summary
      的聚合口径一致，但这里只要"出现最多的那一组"而非全部组合。
    ⛔ 字体须读 eastAsia（_docx_run_font 已兼顾 ascii/eastAsia），中文文档才取得到。
    ⛔ 全部 try 包裹：参考文件可能缺字段或结构异常，任一处失败都回退默认、绝不抛断写入。
    """
    result: dict[str, Any] = {"page": {}, "body": {}}
    try:
        import docx
        d = docx.Document(str(path))
    except Exception:
        return result

    # ── 页面设置 ──
    try:
        s = d.sections[0]

        def _cm(v: Any) -> Any:
            return round(v.cm, 2) if v is not None else None

        result["page"] = {
            "width_cm": _cm(s.page_width), "height_cm": _cm(s.page_height),
            "top_cm": _cm(s.top_margin), "bottom_cm": _cm(s.bottom_margin),
            "left_cm": _cm(s.left_margin), "right_cm": _cm(s.right_margin),
        }
    except Exception:
        pass

    # ── 正文主格式（按出现次数聚合取众数）──
    combos: dict[tuple, int] = {}
    for p in d.paragraphs:
        if not (p.text or "").strip():
            continue
        # 跳过标题段（标题格式不该主导正文）
        try:
            sname = (p.style.name if p.style is not None else "") or "Normal"
        except Exception:
            sname = "Normal"
        low = sname.lower()
        if low.startswith("heading") or sname.startswith("标题"):
            continue
        name, size, _bold = _docx_run_font(p.runs[0]) if p.runs else (None, None, None)
        # 字体名可能是 "ascii/eastAsia" 合并串，写出端要分开 → 这里拆开
        ascii_font = eastasia = None
        if name:
            if "/" in name:
                ascii_font, eastasia = name.split("/", 1)
            else:
                eastasia = name  # 只有 eastAsia 时（中文文档常见）
        fli = lsp_mult = lsp_pt = None
        try:
            pf = p.paragraph_format
            if pf.first_line_indent is not None:
                fli = round(pf.first_line_indent.cm, 2)
            _ls = pf.line_spacing
            if _ls is not None:
                if hasattr(_ls, "pt"):       # Length = 固定行距（磅）
                    lsp_pt = round(float(_ls.pt), 1)
                else:                        # float = 倍数行距
                    lsp_mult = round(float(_ls), 2)
        except Exception:
            pass
        key = (ascii_font, eastasia, size, _align_label(p.alignment), fli, lsp_mult, lsp_pt)
        combos[key] = combos.get(key, 0) + 1

    if combos:
        (ascii_font, eastasia, size, align_lbl, fli, lsp_mult, lsp_pt), _cnt = \
            max(combos.items(), key=lambda kv: kv[1])
        result["body"] = {
            "font": eastasia, "ascii_font": ascii_font, "size_pt": size,
            "align_value": _ALIGN_VALUE.get(align_lbl),
            "first_line_indent_cm": fli,
            "line_spacing": lsp_mult, "line_spacing_pt": lsp_pt,
        }
    return result


# ────────────────────────────── xlsx ──────────────────────────────

def _read_xlsx(path: Path) -> tuple[str, str]:
    import openpyxl

    # data_only=True 取公式的**计算结果**而非公式串（Agent 要的是数据本身）。
    # ⛔ 但它只读 Excel 写入的缓存值：文件若从未被 Excel 打开计算过
    # （程序新建的表、脚本导出的表），缓存为空 → 公式单元格读出来是空白，
    # 会被误当成"原表这格没值"。故同时开 data_only=False 取公式串兜底，
    # 结果缺失时如实显示 `=SUM(B2:B3)` 并标注未计算，绝不静默留空。
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        wb_raw = openpyxl.load_workbook(str(path), data_only=False)
    except Exception:
        wb_raw = None

    out: list[str] = []
    stats: list[str] = []
    n_formula = 0
    for ws in wb.worksheets:
        ws_raw = wb_raw[ws.title] if wb_raw is not None and ws.title in wb_raw.sheetnames else None
        grid: list[list[Any]] = []
        for ri, row in enumerate(ws.iter_rows(values_only=False), start=1):
            vals: list[Any] = []
            for cell in row:
                v = cell.value
                if v is None and ws_raw is not None:
                    rv = ws_raw.cell(row=ri, column=cell.column).value
                    # 结果是空但原始格是公式 → 显示公式串，明确标注未被 Excel 计算过
                    if isinstance(rv, str) and rv.startswith("="):
                        v = f"{rv}（公式未计算）"
                        n_formula += 1
                vals.append(v)
            grid.append(vals)
        rows = [r for r in grid if any((str(c).strip() for c in r if c is not None))]
        stats.append(f"「{ws.title}」{ws.max_row}行×{ws.max_column}列（非空 {len(rows)} 行）")
        if not rows:
            continue
        out.append(f"## 工作表：{ws.title}")
        width = max(len(r) for r in rows)
        for i, r in enumerate(rows):
            cells = [("" if c is None else str(c).replace("\n", " ").strip()) for c in r]
            cells += [""] * (width - len(cells))
            out.append("| " + " | ".join(cells) + " |")
            if i == 0:
                out.append("|" + "---|" * width)
        out.append("")
    detail = ("【表格结构】" + "；".join(stats)) if stats else ""
    if n_formula:
        detail += (f"。⚠️ 有 {n_formula} 个公式单元格缺少计算缓存（文件未被 Excel 打开计算过），"
                   f"已原样显示公式；如需其数值请如实说明无法获得，不要自行推算并当作原表数据。")
    return "\n".join(out), detail


# ────────────────────────────── pptx ──────────────────────────────

def _xml_unescape(s: str) -> str:
    # ⛔ &amp; 必须最后替换，否则 "&amp;lt;" 会被二次解成 "<"
    return (s.replace("&lt;", "<").replace("&gt;", ">")
             .replace("&quot;", '"').replace("&apos;", "'").replace("&amp;", "&"))


def _read_pptx(path: Path) -> tuple[str, str]:
    """pptx 文本提取（zip 兜底，不依赖 python-pptx）。"""
    slides: list[str] = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        names.sort(key=lambda n: int(re.search(r"(\d+)", n).group(1)))
        for n in names:
            xml = z.read(n).decode("utf-8", errors="ignore")
            texts = [_xml_unescape(m).strip() for m in _PPTX_TEXT_RE.findall(xml)]
            texts = [t for t in texts if t]
            idx = int(re.search(r"(\d+)", n).group(1))
            body = "\n".join(f"- {t}" for t in texts) if texts else "（本页无文本，可能为纯图片版式）"
            slides.append(f"## 第 {idx} 页\n{body}")
    detail = (f"【PPTX 结构】共 {len(slides)} 页。"
              f"（未安装 python-pptx，按 OOXML 文本节点提取，故不含版式/母版/动画细节）")
    return "\n\n".join(slides), detail


# ────────────────────────────── pdf ──────────────────────────────

def _read_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError:  # 兼容旧包名
        from PyPDF2 import PdfReader  # type: ignore

    r = PdfReader(str(path))
    n = len(r.pages)
    out: list[str] = []
    empty_pages = 0
    for i, pg in enumerate(r.pages):
        try:
            t = (pg.extract_text() or "").strip()
        except Exception:
            t = ""
        head = f"\n----- 第 {i + 1}/{n} 页 -----"
        if t:
            out.append(head + "\n" + t)
        else:
            empty_pages += 1
            out.append(head + "\n（本页无可提取文本，可能是扫描件/图片型 PDF，需走图像识别）")
    detail = f"【PDF 结构】共 {n} 页"
    if empty_pages:
        detail += f"，其中 {empty_pages} 页无文本层"
    try:
        title = (r.metadata or {}).get("/Title")
        if title:
            detail += f"；标题：{title}"
    except Exception:
        pass
    return "\n".join(out), detail
