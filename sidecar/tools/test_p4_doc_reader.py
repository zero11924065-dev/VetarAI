"""#4（0.4.19）· read_file 真正解析 Office/PDF 文档（而非返回二进制乱码）。

## 缺陷与修复

**缺陷**：read_file 对所有非图片文件一律 `read_bytes().decode("utf-8", errors="replace")`。
.docx/.pptx/.xlsx 是 **zip 压缩包**、.pdf 是**二进制**、旧式 .doc 是 **OLE 复合二进制**，
这样解出来必然是 `PK\\x03\\x04...` 乱码。用户实测：Agent 只能看到二进制垃圾，
完全读不出上传的参考律师函——**这是应用缺口，不是模型能力问题**（图片早有 base64
特例让多模态模型真正看图，文档却没有对应通道）。

**修复**：新增 `sidecar/tools/doc_reader.py`，read_file 对 DOC_EXTS/LEGACY_EXTS
改走解析分支，返回「格式概要 + 正文（含表格）」。

## 覆盖

- T1 docx 正文与表格（含表格必须按文档流顺序、markdown 化）
- T2 格式概要（字体 eastAsia 中文、字号、对齐、首行缩进、行距单位、页面尺寸）
- T3 xlsx 多工作表 + 公式缓存缺失兜底
- T4 pptx 走 zip 兜底（项目未装 python-pptx）+ XML 实体解码
- T5 pdf 文本层提取 + 无文本层如实提示
- T6 乱码防御：任何支持格式都不得出现 PK\\x03\\x04 / 替换符
- T7 契约与预算：ok/content/size 齐备、类型正确、1MB 截断不破契约
- T8 异常路径：损坏 zip、不支持扩展名、文件不存在（不拖垮 read_file）
- T9 旧式 .doc：macOS textutil 转换；不可用时如实提示另存
- T10 含图为主的文档：必须提示图片数量，不得让模型误判"文件是空的"
- T11 回归：纯文本/图片既有行为不变
- T12 变异测试：三处关键修复被撤掉时必须失败（验证本测试真的有效）

⛔ 断言一律用 find()/in，不用 index()——变异模式下子串缺失会抛 ValueError，
   整个套件崩溃并掩盖后续失败（#3/#7 已踩过两次同一个坑）。

运行：./sidecar/.venv/bin/python sidecar/tools/test_p4_doc_reader.py
变异：MUTATE=1|2|3 ./sidecar/.venv/bin/python sidecar/tools/test_p4_doc_reader.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from sidecar.tools import doc_reader as dr
from sidecar.tools.registry import execute, MAX_READ_BYTES

MUTATE = int(os.environ.get("MUTATE", "0"))
_PASS = 0
_FAIL = 0
_FAILURES: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    global _PASS, _FAIL
    if cond:
        _PASS += 1
        print(f"  ✅ {name}")
    else:
        _FAIL += 1
        _FAILURES.append(name)
        print(f"  ❌ {name}  {detail[:400]}")


# ── 变异注入（把修复改回缺陷态，验证测试真的在守这条线）──
# ⛔ 还原必须用**内存备份**，不能用 `git checkout --`：
#   ① 本次修复尚未提交，checkout 会把 registry.py 的 #4 接线直接抹掉；
#   ② doc_reader.py 是新增文件（untracked），git checkout 根本还原不了它。
_BACKUP: dict[str, str] = {}


# ⛔ 变异后必须重取 `execute` 的**当前对象**：本模块开头是
# `from sidecar.tools.registry import execute`，那是按名绑定的函数对象。
# reload(registry) 会在模块命名空间里新建一个 execute，但本模块的名字仍指向旧的
# → 变异 1（撤掉解析分支）会**假通过**。故统一经 _exec() 间接取用。
_EXEC_HOLDER: dict = {}


def _exec(tool: str, args: dict, root: str):
    """调用 read_file 等工具（始终用重载后的最新实现）。"""
    import asyncio
    fn = _EXEC_HOLDER.get("execute")
    if fn is None:
        from sidecar.tools.registry import execute as fn  # type: ignore
        _EXEC_HOLDER["execute"] = fn
    return asyncio.run(fn(tool, args, root))


def _rebind_exec() -> None:
    """重载 registry 后刷新 execute 绑定。"""
    import importlib
    import sidecar.tools.registry as reg
    importlib.reload(reg)
    _EXEC_HOLDER["execute"] = reg.execute


def _read_src(mod) -> str:
    return Path(mod.__file__).read_text(encoding="utf-8")


def _apply_mutation() -> None:
    if not MUTATE:
        return
    import sidecar.tools.registry as reg
    _BACKUP["reg"] = _read_src(reg)
    _BACKUP["dr"] = _read_src(dr)
    if MUTATE == 1:
        # 撤掉解析分支 → 退回按字节解码（原始缺陷态）
        src = _BACKUP["reg"]
        patched = src.replace("if _dr.is_parseable(target):", "if False:")
        assert patched != src, "变异 1 未命中 registry 源码，测试无效"
        Path(reg.__file__).write_text(patched, encoding="utf-8")
    elif MUTATE == 2:
        # 表格/段落不再按文档流顺序取（退回 doc.paragraphs 单独遍历）
        src = _BACKUP["dr"]
        patched = src.replace("for block in _iter_docx_blocks(d):",
                              "for block in list(d.paragraphs):")
        assert patched != src, "变异 2 未命中 doc_reader 源码，测试无效"
        Path(dr.__file__).write_text(patched, encoding="utf-8")
    elif MUTATE == 3:
        # 中文字体不再读 eastAsia（退回只用 font.name）
        src = _BACKUP["dr"]
        patched = src.replace(
            '            if ea:\n                name = ea if (not name or name == ea) else f"{name}/{ea}"',
            '            if ea:\n                pass')
        assert patched != src, "变异 3 未命中 doc_reader 源码，测试无效"
        Path(dr.__file__).write_text(patched, encoding="utf-8")
    else:
        raise SystemExit(f"未知变异编号 {MUTATE}（支持 1|2|3）")


def _restore() -> None:
    """把被变异改写的源文件还原为测试开始前的内容。"""
    if not MUTATE or not _BACKUP:
        return
    import sidecar.tools.registry as reg
    try:
        if "reg" in _BACKUP and _read_src(reg) != _BACKUP["reg"]:
            Path(reg.__file__).write_text(_BACKUP["reg"], encoding="utf-8")
        if "dr" in _BACKUP and _read_src(dr) != _BACKUP["dr"]:
            Path(dr.__file__).write_text(_BACKUP["dr"], encoding="utf-8")
    finally:
        _BACKUP.clear()


# ── 样本构造 ──────────────────────────────────────────────

def _set_run_cjk_font(run, font: str) -> None:
    """把 run 的中文字体写到 w:eastAsia 上（复刻真实中文文档的常态）。

    ⛔ 不能直接 `run._element.rPr.rFonts.set(...)`：python-docx 只在你设过
    font.name 时才创建 rFonts，光设字号时 rPr/rFonts 为 None → AttributeError。
    与 doc_writer._set_cjk_font 同一模式：get_or_add_rPr() 后 find 不到就建。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)          # 中文字体
    # ⛔ 西文必须与中文**不同**（真实中文文档：ascii=Times New Roman、eastAsia=仿宋）。
    # 若两者都设成中文，则只读 font.name(=ascii) 也能拿到中文 → 变异3（丢 eastAsia）
    # 会假通过。分离后，丢 eastAsia 就拿不到中文，断言才守得住这条修复。
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def _make_docx(path: Path, with_table: bool = True, cjk_font: str = "仿宋") -> None:
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)

    t = d.add_paragraph("民事起诉状")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in t.runs:
        r.font.size = Pt(22)
        r.font.bold = True
        _set_run_cjk_font(r, cjk_font)

    p1 = d.add_paragraph("原告：王永斌，男，汉族，住重庆市渝北区。")
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p1.paragraph_format.first_line_indent = Cm(0.99)
    for r in p1.runs:
        r.font.size = Pt(14)
        _set_run_cjk_font(r, cjk_font)

    if with_table:
        # ⛔ 表格必须夹在段落之间，验证"按文档流顺序"而非全部堆到末尾
        d.add_paragraph("诉讼请求如下表所列：")
        tb = d.add_table(rows=3, cols=2)
        tb.cell(0, 0).text = "请求事项"
        tb.cell(0, 1).text = "金额"
        tb.cell(1, 0).text = "返还借款本金"
        tb.cell(1, 1).text = "239900元"
        tb.cell(2, 0).text = "支付资金占用损失"
        tb.cell(2, 1).text = "5972元"
        d.add_paragraph("事实与理由：被告拖欠设备意向金，经多次催告拒不返还。")
    d.save(str(path))


def _make_xlsx(path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "借款明细"
    ws.append(["项目", "金额", "日期"])
    ws.append(["本金", 239900, "2023-02-01"])
    ws.append(["利息", 5972, "2026-06-03"])
    ws.append(["合计", "=SUM(B2:B3)", ""])   # 无计算缓存 → 须原样显示公式
    wb.create_sheet("空工作表")
    wb.save(str(path))


def _make_pptx(path: Path) -> None:
    """手工造 pptx（zip + slide xml）：项目未装 python-pptx，只能走 zip 兜底。"""
    xml = ('<?xml version="1.0"?>'
           '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<p:cSld><p:spTree><p:sp><p:txBody>'
           '<a:p><a:r><a:t>案件汇报</a:t></a:r></a:p>'
           '<a:p><a:r><a:t>本金 239900 &amp; 利息 5972</a:t></a:r></a:p>'
           '</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        z.writestr("ppt/slides/slide2.xml",
                   xml.replace("案件汇报", "第二页").replace("本金", "备注"))
        z.writestr("ppt/slides/slide1.xml", xml)


def _make_pdf(path: Path, with_text: bool = True) -> None:
    """构造结构有效的最小 PDF。

    ⛔ 踩坑记录：首版把 bytes 经 f-string 拼进对象体，写入的是 b'...' 的 repr，
    生成的 PDF 结构非法 → pypdf 提取不到文本，差点误判成"pypdf 不行"。
    验证解析能力前，先确认样本自身有效（字节拼接，不走 f-string）。
    """
    body_text = (b'BT /F1 16 Tf 72 720 Td (Wang Yongbin Loan 239900) Tj ET\n'
                 b'BT /F1 16 Tf 72 690 Td (Plaintiff: Wang Yongbin) Tj ET')
    content = body_text if with_text else b'0 0 612 792 re f'
    objs = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
        b'/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',
        b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    pdf = bytearray(b'%PDF-1.4\n')
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(pdf))
        pdf += str(i).encode() + b' 0 obj\n' + body + b'\nendobj\n'
    xref = len(pdf)
    pdf += b'xref\n0 ' + str(len(objs) + 1).encode() + b'\n0000000000 65535 f \n'
    for off in offsets:
        pdf += (f'{off:010d} 00000 n \n').encode()
    pdf += (b'trailer\n<< /Size ' + str(len(objs) + 1).encode()
            + b' /Root 1 0 R >>\nstartxref\n' + str(xref).encode() + b'\n%%EOF\n')
    path.write_bytes(bytes(pdf))


def _make_docx_with_images(path: Path, n_images: int = 3) -> None:
    """含图为主、几乎无文字的 docx（复刻用户实测的 7.6MB 证据文档场景）。"""
    import docx
    d = docx.Document()
    d.add_paragraph("证据材料")
    # 直接往包里塞图片字节（比造真图简单，且 docx 只按 word/media/ 统计）
    d.save(str(path))
    with zipfile.ZipFile(path, "a") as z:
        for i in range(n_images):
            z.writestr(f"word/media/image{i + 1}.png", b"\x89PNG\r\n\x1a\n" + b"0" * 5000)


# ── 测试 ──────────────────────────────────────────────────

def t1_docx_body_and_table(tmp: Path) -> None:
    print("\nT1 docx：正文 + 表格按文档流顺序提取")
    p = tmp / "起诉状.docx"
    _make_docx(p)
    body, detail = dr._read_docx(p)
    check("正文含原告信息", "原告：王永斌" in body, body[:200])
    check("正文含事实与理由", "事实与理由" in body, body[:200])
    check("标题保留", "民事起诉状" in body, body[:200])
    # ⛔ 表格内容必须提取到（此前只遍历 paragraphs 会全丢）
    check("表格单元格已提取", "返还借款本金" in body and "239900元" in body, body[:400])
    check("表格已 markdown 化（含分隔行）", body.find("|---") >= 0, body[:400])
    # ⛔ 顺序：表格必须夹在两段之间，不是全部堆到末尾
    i_lead = body.find("诉讼请求如下表所列")
    i_cell = body.find("返还借款本金")
    i_tail = body.find("事实与理由")
    check("表格位于引导段与结尾段之间（文档流顺序）",
          0 <= i_lead < i_cell < i_tail,
          f"lead={i_lead} cell={i_cell} tail={i_tail}")


def t2_format_summary(tmp: Path) -> None:
    print("\nT2 格式概要：中文字体/字号/对齐/缩进/行距/页面")
    p = tmp / "格式.docx"
    _make_docx(p, cjk_font="仿宋")
    _, detail = dr._read_docx(p)
    check("概要段存在", "【格式概要】" in detail, detail[:300])
    # ⛔ 中文字体挂在 w:eastAsia 上；只读 font.name 会恒为 None（真实中文文档等于没提取）
    check("中文字体（eastAsia）已提取", "仿宋" in detail, detail[:500])
    check("字号已提取", "字号=22" in detail or "字号=14" in detail, detail[:500])
    check("对齐已转人话", "居中" in detail and "两端对齐" in detail, detail[:500])
    check("首行缩进已提取", "首行缩进=0.99" in detail, detail[:500])
    check("页面尺寸与方向", "纵向" in detail and "21.0" in detail, detail[:300])
    check("页边距已提取", "页边距" in detail, detail[:300])
    check("表格数量已统计", "表格：1 个" in detail, detail[:500])
    # ⛔ 行距不得是 EMU 原值（22磅 = 279400 EMU，直出等于给模型看鬼数字）
    check("行距不是 EMU 原值", "279400" not in detail, detail[:600])


def t3_xlsx(tmp: Path) -> None:
    print("\nT3 xlsx：多工作表 + 公式缓存缺失兜底")
    p = tmp / "借款.xlsx"
    _make_xlsx(p)
    body, detail = dr._read_xlsx(p)
    check("数据行已提取", "239900" in body and "5972" in body, body[:300])
    check("表头已提取", "项目" in body and "金额" in body, body[:300])
    check("工作表名已列出", "借款明细" in body, body[:300])
    check("结构统计含两个表", "空工作表" in detail, detail[:300])
    # ⛔ data_only=True 只读缓存值；程序新建的表没有缓存 → 公式格会是空白，
    # 会被误当成"原表这格没值"。必须原样显示公式并标注。
    check("公式未静默留空（显示公式串）", "=SUM(B2:B3)" in body, body[:400])
    check("公式缺失已明确标注", "公式未计算" in body, body[:400])
    check("概要提示了公式数量", "1 个公式单元格" in detail, detail[:400])


def t4_pptx(tmp: Path) -> None:
    print("\nT4 pptx：zip 兜底提取 + XML 实体解码")
    p = tmp / "汇报.pptx"
    _make_pptx(p)
    body, detail = dr._read_pptx(p)
    check("文本已提取", "案件汇报" in body, body[:300])
    # ⛔ &amp; 必须解码为 &；且替换顺序错了会把 &amp;lt; 二次解成 <
    check("XML 实体已解码", "本金 239900 & 利息 5972" in body, body[:400])
    check("实体未残留", "&amp;" not in body, body[:400])
    check("按页码排序（slide1 在 slide2 前）",
          body.find("案件汇报") < body.find("第二页"), body[:400])
    check("页数统计正确", "共 2 页" in detail, detail[:200])
    check("如实说明未装 python-pptx", "python-pptx" in detail, detail[:300])


def t5_pdf(tmp: Path) -> None:
    print("\nT5 pdf：文本层提取 + 无文本层如实提示")
    p = tmp / "有效.pdf"
    _make_pdf(p, with_text=True)
    body, detail = dr._read_pdf(p)
    check("PDF 文本已提取", "Wang Yongbin Loan 239900" in body, body[:300])
    check("页数统计", "共 1 页" in detail, detail[:200])
    check("分页标记存在", "第 1/1 页" in body, body[:300])

    p2 = tmp / "扫描件.pdf"
    _make_pdf(p2, with_text=False)
    body2, detail2 = dr._read_pdf(p2)
    check("无文本层如实提示（不谎称提取到内容）", "无可提取文本" in body2, body2[:300])
    check("提示指向图像识别", "图像识别" in body2, body2[:300])
    check("结构概要标注无文本页数", "无文本层" in detail2, detail2[:200])


def t6_no_garbage(tmp: Path) -> None:
    print("\nT6 乱码防御：支持格式一律不得返回二进制垃圾")
    samples = []
    for name, maker in (("a.docx", lambda p: _make_docx(p)),
                        ("b.xlsx", _make_xlsx),
                        ("c.pptx", _make_pptx),
                        ("d.pdf", lambda p: _make_pdf(p))):
        p = tmp / name
        maker(p)
        samples.append(p)
    # ⛔ 对照组：证明这些文件按旧逻辑（字节解码）确实会产生乱码
    old_style = samples[0].read_bytes()[:200].decode("utf-8", errors="replace")
    check("对照组：旧逻辑对 docx 产生乱码（证明缺陷真实存在）",
          "\ufffd" in old_style or "PK" in old_style, repr(old_style[:60]))
    for p in samples:
        r = dr.extract(p, MAX_READ_BYTES)
        c = r["content"]
        head = c[:400]
        check(f"{p.suffix} 无 zip 魔数 PK\\x03\\x04",
              c.find("PK\x03\x04") < 0, repr(head[:80]))
        check(f"{p.suffix} 无 Unicode 替换符（乱码特征）",
              "\ufffd" not in c[:2000], repr(head[:80]))
        check(f"{p.suffix} 解析成功且有实质内容",
              r["ok"] is True and len(c.strip()) > 20, f"len={len(c)}")


def t7_contract_and_budget(tmp: Path) -> None:
    print("\nT7 返回契约与字节预算")
    p = tmp / "契约.docx"
    _make_docx(p)
    r = _exec("read_file", {"path": str(p)}, str(tmp))
    check("required 字段齐备", all(k in r for k in ("ok", "content", "size")), str(list(r)))
    check("类型正确",
          isinstance(r["ok"], bool) and isinstance(r["content"], str)
          and isinstance(r["size"], int), str({k: type(v).__name__ for k, v in r.items()}))
    check("ok=True", r["ok"] is True, str(r)[:200])
    check("size 为文件真实字节数", r["size"] == p.stat().st_size, str(r.get("size")))
    check("content 非空", len(r["content"]) > 50, str(r.get("content"))[:120])
    # ⛔⛔ 端到端强断言（专守变异1：撤掉 registry 的解析接线）。
    # 上面 len>50 太弱——乱码也超 50 字。必须断言【经 registry】读 docx 得到的是
    # **解析后的真实正文**而非字节解码乱码：既查真实内容存在，又查乱码特征不存在。
    # 变异1 把 `if _dr.is_parseable(target):` 改成 `if False:` → 落到 read_bytes().decode
    # → content 变 PK\x03\x04 乱码 → 这三条至少一条失败。
    _rc = r["content"]
    check("【经 registry】docx 含真实正文（非乱码）",
          "原告：王永斌" in _rc and "239900元" in _rc, repr(_rc[:120]))
    check("【经 registry】docx 无 zip 魔数 PK",
          _rc.find("PK\x03\x04") < 0, repr(_rc[:80]))
    check("【经 registry】docx 无 Unicode 替换符",
          "\ufffd" not in _rc[:2000], repr(_rc[:80]))
    check("【经 registry】docx 带解析头部标记",
          "已解析" in _rc, repr(_rc[:80]))

    # 预算收紧 → 必须截断且总长受控，不得超预算
    small = dr.extract(p, 200)
    check("小预算触发截断标记", small.get("truncated") is True, str(small.get("truncated")))
    check("截断后长度受控", len(small["content"].encode("utf-8")) <= 200 + 60,
          str(len(small["content"].encode("utf-8"))))
    check("截断有明确说明", "已截断" in small["content"], small["content"][-80:])


def t8_error_paths(tmp: Path) -> None:
    print("\nT8 异常路径：不得拖垮 read_file")
    # 扩展名是 docx 但内容是垃圾 → BadZipFile 必须被捕获
    # ⛔ bytes 字面量不能含中文（SyntaxError），中文须经 encode
    bad = tmp / "坏文件.docx"
    bad.write_bytes("这不是一个真正的 docx，只是扩展名叫 docx".encode("utf-8"))
    r = dr.extract(bad, MAX_READ_BYTES)
    check("损坏文件 ok=False 而非抛异常", r["ok"] is False, str(r)[:200])
    check("损坏原因如实说明", "不是有效的" in r["content"] or "出错" in r["content"],
          r["content"][:200])
    check("明确要求不得编造", "不要编造" in r["content"], r["content"][:250])

    # 不支持的扩展名走 unsupported_ext 分支
    r2 = dr.extract(Path("/tmp/不存在.xyz"), MAX_READ_BYTES)
    check("不支持扩展名返回失败说明", r2["ok"] is False, str(r2)[:160])

    # read_file 端到端：文件不存在仍走原有错误路径（未回归）
    r3 = _exec("read_file", {"path": str(tmp / "没有.docx")}, str(tmp))
    check("文件不存在仍报 not_a_file（未回归）",
          r3.get("ok") is False and "not_a_file" in str(r3.get("error", "")), str(r3)[:200])


def t9_legacy_doc(tmp: Path) -> None:
    print("\nT9 旧式 .doc：textutil 转换 / 不可用时如实提示")
    check(".doc 纳入可解析范围", dr.is_parseable(Path("a.doc")) is True)
    check(".xls/.ppt 同样纳入", dr.is_parseable(Path("a.xls")) and dr.is_parseable(Path("a.ppt")))
    check(".txt 不走解析分支", dr.is_parseable(Path("a.txt")) is False)

    import shutil as _sh
    has_textutil = bool(_sh.which("textutil"))
    # 用 textutil 造一个真实 .doc（系统能力可用时才有意义）
    src = tmp / "旧文档.txt"
    src.write_text("借据\n今借到王永斌人民币贰拾叁万玖仟玖佰元整（¥239900.00）。\n借款人：张元\n",
                   encoding="utf-8")
    docp = tmp / "旧文档.doc"
    if has_textutil:
        import subprocess
        subprocess.run(["textutil", "-convert", "doc", str(src), "-output", str(docp)],
                       capture_output=True, timeout=60)
    if docp.exists() and docp.stat().st_size > 0:
        r = dr.extract(docp, MAX_READ_BYTES)
        check("textutil 转换后成功解析 .doc", r["ok"] is True, str(r)[:300])
        check(".doc 中文与金额完好", "239900" in r["content"], r["content"][:400])
        check(".doc 正文含借据内容", "借据" in r["content"] or "王永斌" in r["content"],
              r["content"][:400])
        # ⛔ 临时目录必须清理（转换产物是中间件，不该留在磁盘）
        leftovers = list(Path(tempfile.gettempdir()).glob("docreader_*"))
        check("转换临时目录已清理", len(leftovers) == 0, str(leftovers[:3]))
    else:
        r = dr.extract(docp, MAX_READ_BYTES)
        check("无 textutil 时如实提示另存", "另存为" in r["content"], r["content"][:300])

    # .xls/.ppt：textutil 不支持 → 必须如实提示，不得假装解析成功
    fake_xls = tmp / "表格.xls"
    fake_xls.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 500)  # OLE 魔数
    r2 = dr.extract(fake_xls, MAX_READ_BYTES)
    check(".xls 不被谎称解析成功", r2["ok"] is False, str(r2)[:200])
    check(".xls 提示另存为 .xlsx", "xlsx" in r2["content"], r2["content"][:300])


def t10_image_heavy_doc(tmp: Path) -> None:
    print("\nT10 含图为主的文档：不得让模型误判'文件是空的'")
    p = tmp / "证据.docx"
    _make_docx_with_images(p, n_images=25)
    r = dr.extract(p, MAX_READ_BYTES)
    c = r["content"]
    check("如实统计内嵌图片数", "25 张" in c, c[:400])
    check("提示图片内容需走图像识别", "图像识别" in c, c[:500])
    check("明确禁止臆测图片内容", "不得据本文本臆测" in c or "臆测" in c, c[:600])
    # ⛔ 文本极少时也要给出可判断信息，不能只回几十字让模型以为文件空了
    check("返回内容含实质提示信息", len(c.strip()) > 80, f"len={len(c)}")


def t11_regression(tmp: Path) -> None:
    print("\nT11 回归：纯文本与图片既有行为不变")
    txt = tmp / "说明.txt"
    txt.write_text("这是纯文本文件，内容不应被文档解析影响。239900", encoding="utf-8")
    r = _exec("read_file", {"path": str(txt)}, str(tmp))
    check("txt 内容原样返回", "这是纯文本文件" in r.get("content", ""), str(r)[:200])
    check("txt 无解析头部标记", "已解析" not in r.get("content", ""), str(r)[:120])
    check("txt ok=True", r.get("ok") is True, str(r)[:160])

    img = tmp / "图.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"fakeimagedata" * 50)
    r2 = _exec("read_file", {"path": str(img)}, str(tmp))
    check("图片仍走 base64 特例", r2.get("_kind") == "image", str(list(r2))[:200])
    check("图片仍带 base64", bool(r2.get("image_base64")), str(r2)[:160])

    # 大文本软提示（B2/0.4.8）仍生效
    big = tmp / "大文本.txt"
    big.write_text("啊" * 300000, encoding="utf-8")
    r3 = _exec("read_file", {"path": str(big)}, str(tmp))
    check("大文本软提示未回归", "大文件提示" in r3.get("content", ""), str(r3)[:160])


def t12_real_user_files() -> None:
    print("\nT12 真实用户文档（存在才跑，不存在则跳过不算失败）")
    real = [
        "/Users/vetar/Desktop/王永斌证据材料/律师函（方贤君）.docx",
        "/Users/vetar/Desktop/王永斌证据材料/律师函.docx",
        "/Users/vetar/Desktop/王永斌证据材料/证据目录及说明（王永斌）.docx",
        "/Users/vetar/Desktop/王永斌证据材料/证据（王永斌）.docx",
    ]
    ran = 0
    for f in real:
        p = Path(f)
        if not p.exists():
            continue
        ran += 1
        r = dr.extract(p, MAX_READ_BYTES)
        c = r["content"]
        check(f"{p.name} 解析成功", r["ok"] is True, str(r)[:200])
        check(f"{p.name} 无乱码",
              c.find("PK\x03\x04") < 0 and "\ufffd" not in c[:2000], repr(c[:80]))
        check(f"{p.name} 提取到实质内容", len(c.strip()) > 60, f"len={len(c)}")
    if ran == 0:
        print("  ⏭ 真实文档不在本机，跳过（不计失败）")
    else:
        print(f"  （已验证 {ran} 个真实文档）")


def main() -> int:
    print("=" * 72)
    print(f"#4 read_file 文档解析测试  |  变异模式 = {MUTATE}")
    print("=" * 72)
    _apply_mutation()
    # 变异改写源码后须重新导入，否则测的是旧模块；并刷新 execute 绑定（见 _exec 注释）
    if MUTATE:
        import importlib
        importlib.reload(dr)
        _rebind_exec()
    try:
        with tempfile.TemporaryDirectory(prefix="p4_") as td:
            tmp = Path(td)
            t1_docx_body_and_table(tmp)
            t2_format_summary(tmp)
            t3_xlsx(tmp)
            t4_pptx(tmp)
            t5_pdf(tmp)
            t6_no_garbage(tmp)
            t7_contract_and_budget(tmp)
            t8_error_paths(tmp)
            t9_legacy_doc(tmp)
            t10_image_heavy_doc(tmp)
            t11_regression(tmp)
        t12_real_user_files()
    finally:
        _restore()
        if MUTATE:
            import importlib
            importlib.reload(dr)
            _rebind_exec()

    print("\n" + "=" * 72)
    print(f"结果：{_PASS} 通过 / {_FAIL} 失败")
    if _FAILURES:
        print("失败项：" + "；".join(_FAILURES))
    if MUTATE and _FAIL == 0:
        print(f"⛔ 变异 {MUTATE} 未被抓住 —— 本测试对该修复无效，断言需加强")
    print("=" * 72)
    return 1 if _FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
