# VetarAI - Local-first multi-agent orchestration application
# Copyright (C) 2026 zero11924065-dev
# GPL-3.0-or-later（见仓库 LICENSE）
"""A10（0.4.18）专项：知识仓库索引本地文档（扩非 md）+ 拉模式铁律 + 数据安全。

═══ 治的是什么 ═══
rebuild_index() 原本只 glob("*.md")，用户把 docx/pdf 丢进知识目录索引不到、agent 检索不到。
现扩为复用 attachments/parser.py 解析多种文档（与聊天附件、工作流 file_read 同一解析链）。

═══ 关键约束（勿凭直觉改）═══
1. ⛔ **拉模式铁律**：只是把文本索引进 FTS/向量供 search_knowledge **检索**，
   ⛔ 绝不自动注入上下文（防重蹈 0.4.8"主 Agent 自读 90KB PDF 跑 20 分钟"覆辙）。
2. ⛔ **数据安全**：source='file' 是用户自己的原始文件，删条目时**绝不能 unlink**
   （否则用户在面板删一条索引就永久删掉自己的文档）。
3. ⛔ **CHECK 约束迁移**：source 原 CHECK IN ('chat','manual')，写 'file' 会 IntegrityError；
   SQLite 的 CHECK 不能 ALTER → 必须检测旧表并重建迁移（保留现有行）。
4. ⛔ **get_entry 对二进制不崩**：正文读取原用 read_text+_parse_md，对 docx/pdf 抛
   UnicodeDecodeError（非 OSError，原 except 捕不到）→ 点开导入条目详情即崩溃。
5. ⛔ **稳定 id**：外部文件无 frontmatter id，用路径派生 uuid5，重建多次不漂移。
6. ⛔ **端点走 executor**：rebuild_index 解析多文档变 CPU 密集，async 端点直接调会阻塞事件循环。
7. ⛔ **格式清单从 parser 推导**，不另写一份（防 C3 刚清理的双源漂移重犯）。

运行：.venv/bin/python -m sidecar.knowledge.test_a10_warehouse_docs
（⛔ 勿裸跑：走 scripts/run_backend_tests.py 隔离 runner）
"""
import asyncio
import io
import sqlite3
import sys
import tempfile
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

PASS, FAIL = 0, 0
FAILURES = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"PASS  {name}")
    else:
        FAIL += 1
        FAILURES.append(name)
        print(f"FAIL  {name}  {detail}")


def _mk_docx(marker: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph(f"标记内容 {marker} 第一段")
    d.add_paragraph("第二段：关于知识仓库导入文档的说明。")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def _mk_pdf(marker: str) -> bytes:
    """零依赖手写最小含文本 PDF（同 C4 测试，不依赖未安装的 reportlab）。"""
    content = f"BT /F1 24 Tf 72 720 Td ({marker}) Tj ET".encode("ascii")
    pdf = b"%PDF-1.4\n"
    offs = []

    def add(num, body):
        nonlocal pdf
        offs.append(len(pdf))
        pdf += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    add(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    add(2, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    add(3, b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
           b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>")
    add(4, b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    add(5, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    xref = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n"
    for o in offs:
        pdf += f"{o:010d} 00000 n \n".encode()
    pdf += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF"
    return pdf


def main():
    import sidecar.knowledge.warehouse as wh
    import sidecar.storage.store as store
    from sidecar.attachments import parser as P

    # 隔离：数据根 + 索引库 + 项目库全部指向临时目录
    tmp = Path(tempfile.mkdtemp(prefix="a10_"))
    wh._DATA_ROOT_OVERRIDE = tmp
    wh._INDEX_DB_PATH = tmp / "index.db"
    proj_root = tmp / "projects"; proj_root.mkdir(parents=True, exist_ok=True)
    store.PROJECTS_ROOT = proj_root
    store._GDB = proj_root / "_global.db"
    work_dir = tmp / "proj_work"; work_dir.mkdir(parents=True, exist_ok=True)
    pid = store.create_project("A10测试项目", work_dir)
    # ⛔ 用 project_knowledge_dir()（内部会 mkdir）取目录，而非手动拼路径——
    #    真实使用中目录由该函数创建；手动拼路径不 mkdir 会让后续 write_bytes 崩（本测试踩过）。
    kdir = wh.project_knowledge_dir(pid)
    assert kdir is not None and kdir.is_dir(), f"项目知识库目录未就绪：{kdir}"
    gdir = wh.global_knowledge_dir()

    # ── T1 格式清单从 parser 推导（不双源漂移）──
    expect = frozenset(P.SUPPORTED_EXTS - P.IMAGE_EXTS - {".md"})
    check("T1a INDEXABLE_DOC_EXTS == parser 推导值", wh.INDEXABLE_DOC_EXTS == expect,
          f"{sorted(wh.INDEXABLE_DOC_EXTS ^ expect)}")
    check("T1b 含 .docx/.pdf/.doc/.xlsx/.pptx",
          all(e in wh.INDEXABLE_DOC_EXTS for e in (".docx", ".pdf", ".doc", ".xlsx", ".pptx")))
    check("T1c 图片不在可索引清单（无语义文本）", not (wh.INDEXABLE_DOC_EXTS & P.IMAGE_EXTS))
    check("T1d .md 不在（走 frontmatter 路径）", ".md" not in wh.INDEXABLE_DOC_EXTS)
    check("T1e SOURCE_IMPORTED_FILE == 'file'", wh.SOURCE_IMPORTED_FILE == "file")

    # ── T2 CHECK 约束：旧库迁移（source 原仅 chat/manual）──
    old_db = tmp / "legacy.db"
    con = sqlite3.connect(str(old_db))
    con.executescript("""
        CREATE TABLE knowledge_entries (
            id TEXT PRIMARY KEY, title TEXT NOT NULL,
            scope TEXT NOT NULL CHECK(scope IN ('project','global')),
            project_id TEXT, category TEXT, keywords TEXT,
            source TEXT NOT NULL DEFAULT 'chat' CHECK(source IN ('chat','manual')),
            file_path TEXT NOT NULL, created_at TEXT NOT NULL);
        INSERT INTO knowledge_entries VALUES
            ('legacy1','旧条目','global','','','[]','chat','/x/y.md','2026-01-01 00:00:00');
    """)
    con.commit(); con.close()
    saved = wh._INDEX_DB_PATH
    wh._INDEX_DB_PATH = old_db
    conn = wh._iconn()                         # 触发 _ensure_schema 迁移
    schema = conn.execute("SELECT sql FROM sqlite_master WHERE name='knowledge_entries'").fetchone()[0]
    conn.close()
    check("T2a 迁移后 CHECK 含 'file'", "'file'" in schema, schema[:120])
    con = sqlite3.connect(str(old_db))
    kept = con.execute("SELECT COUNT(*) FROM knowledge_entries WHERE id='legacy1'").fetchone()[0]
    check("T2b 迁移保留旧数据（不丢用户条目）", kept == 1, str(kept))
    # 迁移后能写入 file 条目（旧约束会在此 IntegrityError）
    try:
        con.execute("INSERT INTO knowledge_entries VALUES "
                    "('f1','t','global','','','[]','file','/x/z.docx','2026-01-01 00:00:00')")
        con.commit(); ok = True; err = ""
    except sqlite3.IntegrityError as e:
        ok = False; err = str(e)
    finally:
        con.close()
    check("T2c 迁移后可写 source='file'（旧约束会 IntegrityError）", ok, err)
    # 幂等：再次 _iconn 不崩
    try:
        wh._iconn().close(); idem = True
    except Exception as e:
        idem = False; print("  幂等失败:", e)
    check("T2d _ensure_schema 幂等（重复连接不崩）", idem)
    wh._INDEX_DB_PATH = saved

    # ── T3 多格式索引：docx/pdf/txt 丢进知识目录 → rebuild 索引到 ──
    (kdir / "合同.docx").write_bytes(_mk_docx("合同甲"))
    (kdir / "判决.pdf").write_bytes(_mk_pdf("JUDGEPDFYI"))
    (kdir / "笔记.txt").write_text("纯文本知识丙：会议纪要要点。", encoding="utf-8")
    (kdir / "数据.csv").write_text("列甲,列乙\n1,2", encoding="utf-8")
    n = wh.rebuild_index()
    proj_entries = wh.list_entries("project", pid)
    titles = {e["title"] for e in proj_entries}
    check("T3a rebuild 索引到非 md 文档（≥4 条）", n >= 4, f"n={n}")
    check("T3b docx 被索引", "合同" in titles, str(titles))
    check("T3c pdf 被索引", "判决" in titles, str(titles))
    check("T3d txt 被索引", "笔记" in titles, str(titles))
    check("T3e csv 被索引", "数据" in titles, str(titles))
    # 所有导入条目 source=file
    check("T3f 导入条目 source 均为 'file'",
          all(e["source"] == "file" for e in proj_entries), str([e["source"] for e in proj_entries]))

    # ── T4 拉模式铁律：可检索（search 命中正文），而非自动注入 ──
    hits = wh.search_entries("知识仓库导入", "project", pid)
    check("T4a docx 正文可被关键词检索（FTS 命中）",
          any("合同" in h.get("title", "") for h in hits), str([h.get("title") for h in hits]))
    hits2 = wh.search_entries("JUDGEPDFYI", "project", pid)
    check("T4b pdf 解析文本可被检索", any("判决" in h.get("title", "") for h in hits2),
          str([h.get("title") for h in hits2]))
    # ⛔ 拉模式：rebuild_index 不得有任何"注入上下文"副作用（源码断言）
    src = Path(wh.__file__).read_text(encoding="utf-8")
    check("T4c rebuild_index 不含自动注入上下文的调用",
          "inject" not in src.split("def rebuild_index")[1].split("\ndef ")[0].lower())

    # ── T5 get_entry 对二进制不崩溃（我修的回归）──
    docx_entry = next((e for e in proj_entries if e["title"] == "合同"), None)
    check("T5a 找到 docx 条目", docx_entry is not None)
    if docx_entry:
        try:
            got = wh.get_entry(docx_entry["id"])
            crash = False
        except UnicodeDecodeError:
            got = None; crash = True
        except Exception as e:
            got = None; crash = True; print("  其他异常:", type(e).__name__, e)
        check("T5b get_entry(docx) 不抛 UnicodeDecodeError", not crash)
        check("T5c get_entry(docx) 返回解析正文（非空、非乱码）",
              got is not None and "合同甲" in got.get("body", ""), (got or {}).get("body", "")[:60])

    # ── T6 稳定 id：重建两次不漂移 ──
    ids1 = {e["title"]: e["id"] for e in wh.list_entries("project", pid)}
    wh.rebuild_index()
    ids2 = {e["title"]: e["id"] for e in wh.list_entries("project", pid)}
    check("T6a 重建两次 id 不漂移（uuid5 路径派生）", ids1 == ids2,
          f"{ids1.get('合同')} vs {ids2.get('合同')}")
    # id 确实是 uuid5 派生（确定性）
    fp = (kdir / "合同.docx").resolve()
    expect_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"knowledge-file://{fp}"))
    check("T6b id == uuid5(路径)（确定性可复算）", ids2.get("合同") == expect_id,
          f"{ids2.get('合同')} vs {expect_id}")

    # ── T7 数据安全：删 file 条目不 unlink 用户文件 ──
    docx_path = kdir / "合同.docx"
    check("T7a 删前用户文件存在", docx_path.exists())
    if docx_entry:
        wh.delete_entry(docx_entry["id"])
    check("T7b 删条目后索引已移除",
          not any(e["title"] == "合同" for e in wh.list_entries("project", pid)))
    check("T7c ⛔ 删 file 条目**不删除用户原始文件**（数据安全守卫）",
          docx_path.exists(), "用户 docx 被误删！")
    # 对照：manual 条目删时仍删 .md（本模块生成的，语义正确）
    e_md = wh.add_entry("project", pid, "手动条目", "内容丁", keywords=["丁"])
    md_path = Path(e_md["file_path"])
    check("T7d manual 条目文件已生成", md_path.exists())
    wh.delete_entry(e_md["id"])
    check("T7e 删 manual 条目**仍删 .md**（本模块产物，语义正确）", not md_path.exists())

    # ── T8 边界：损坏/空/超大文件跳过，不索引垃圾、不崩 ──
    (kdir / "损坏.pdf").write_bytes(b"not a pdf at all")
    (kdir / "空.docx").write_bytes(b"")
    n2 = wh.rebuild_index()
    titles2 = {e["title"] for e in wh.list_entries("project", pid)}
    check("T8a 损坏 pdf 不被索引", "损坏" not in titles2, str(titles2))
    check("T8b 空文件不被索引", "空" not in titles2, str(titles2))
    check("T8c 损坏/空文件不致 rebuild 崩溃（仍索引到其他）",
          "判决" in titles2 and "笔记" in titles2, str(titles2))

    # ── T9 .md 路径行为不变（回归保护）──
    e_md2 = wh.add_entry("project", pid, "原md条目", "md 正文戊", keywords=["戊"])
    wh.rebuild_index()
    got_md = wh.get_entry(e_md2["id"])
    check("T9a .md 条目仍可索引并读正文", got_md and "md 正文戊" in got_md.get("body", ""),
          (got_md or {}).get("body", "")[:60])
    check("T9b .md 条目 source 不是 file（manual/chat 路径未变）",
          got_md and got_md["source"] != "file", (got_md or {}).get("source"))

    # ── T10 端点走 executor（源码断言：阻塞风险防护）──
    # ⛔⛔ 原版 `"run_in_executor" in eseg` 是**无效断言**（变异测试暴露）：eseg 含 docstring，
    #    而 docstring 里正好写了"必须 run_in_executor"几个字 → 把实际调用改成直接
    #    `_wh.rebuild_index()` 后，docstring 的词仍让断言通过（变异4 全绿）。
    #    ✅ 改为正则锚定**实际调用形态** `run_in_executor(None, _wh.rebuild_index)`：
    #    docstring 的中文说明不含这个精确形态，故不会被污染。
    import re as _re
    app_src = (Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8")
    eseg = app_src.split("async def api_knowledge_rebuild")[1].split("\n@app.")[0]
    check("T10 重建端点用 run_in_executor 实际调用 rebuild_index（删 executor 即失败）",
          _re.search(r"run_in_executor\(\s*None\s*,\s*_wh\.rebuild_index\s*\)", eseg) is not None,
          eseg[:240])

    print(f"\n===== A10 专项: PASS={PASS} FAIL={FAIL} =====")
    if FAILURES:
        for f in FAILURES:
            print("  ✗", f)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
