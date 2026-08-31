"""TS-113 M7 体验与契约增强专项单测（mock + 临时目录，不碰真实 ~/.subagent）。
覆盖：
- 交卷契约扩容（1000 字全文/超长落盘+截断标注/兜底打包全文）
- 默认导出目录解析（空→项目工作目录 / 配置值→配置目录 / 不可用回退）
- 会话 Markdown 导出（文件存在 + 含消息 + 含工具步骤）
- 附件内置解析器（文本/CSV/图片标注/二进制/损坏降级；PDF/Word/Excel 可选）
- 新配置项校验（default_export_dir / vision_parse_attachments）
venv 内 python test_m7.py 直接跑（需 PYTHONPATH）。只输出 PASS/FAIL 摘要。
"""
import asyncio
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

PASS, FAIL = 0, 0
FAILURES = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"PASS  {name}")
    else:
        FAIL += 1
        FAILURES.append(name)
        print(f"FAIL  {name}  {detail}")


def main():
    TMP = Path(tempfile.mkdtemp(prefix="m7_"))

    # ── 隔离配置与存储 ──
    import sidecar.config as cfgmod
    cfgmod.get_config_path = lambda: TMP / "config.json"
    cfgmod._MEM = {}
    from sidecar.config import reload_config, get_config

    import sidecar.storage.store as store
    store.PROJECTS_ROOT = TMP / "projects"
    store._GDB = TMP / "projects" / "_global.db"
    store.PROJECTS_ROOT.mkdir(parents=True, exist_ok=True)

    # ══ 1. 新配置项校验 ══
    reload_config({"default_export_dir": ""})
    check("1a default_export_dir 空合法", get_config()["default_export_dir"] == "")
    try:
        reload_config({"default_export_dir": 123})
        check("1b default_export_dir 非字符串拒绝", False, "未抛错")
    except ValueError:
        check("1b default_export_dir 非字符串拒绝", True)
    reload_config({"vision_parse_attachments": True})
    check("1c vision_parse_attachments bool 可保存", get_config()["vision_parse_attachments"] is True)
    try:
        reload_config({"vision_parse_attachments": "yes"})
        check("1d vision_parse_attachments 非 bool 拒绝", False, "未抛错")
    except ValueError:
        check("1d vision_parse_attachments 非 bool 拒绝", True)
    reload_config({"vision_parse_attachments": False, "default_export_dir": ""})

    # ══ 1e~1g. checkpoint-047：插件/技能逐项开关（替代 046 的全局总闸）══
    # 隔离数据目录（技能目录 = data_root()/skills，避免碰真实 ~/.subagent）
    cfgmod.data_root = lambda: TMP / "dataroot"
    (TMP / "dataroot").mkdir(parents=True, exist_ok=True)

    from sidecar.agent_engine.loop import tools_spec
    names = [s["function"]["name"] for s in tools_spec()]
    check("1e tools_spec 恒含 read_skill（逐项语义：工具常驻，禁用项路由拒绝）",
          "read_skill" in names, str(names))
    # 技能清单只含启用项（逐项开关的核心生效点）
    import sidecar.skills_mgr.manager as _skm
    _sd = _skm.skills_root() / "技能甲"
    _sd.mkdir(parents=True, exist_ok=True)
    (_sd / "SKILL.md").write_text("---\nname: 技能甲\ndescription: 测试\nenabled: true\n---\n正文", encoding="utf-8")
    _sd2 = _skm.skills_root() / "技能乙"
    _sd2.mkdir(parents=True, exist_ok=True)
    (_sd2 / "SKILL.md").write_text("---\nname: 技能乙\ndescription: 测试2\nenabled: false\n---\n正文", encoding="utf-8")
    _lst_text = _skm.build_skills_list_text()
    check("1f 技能清单只含启用项（禁用项不注入提示词）",
          "技能甲" in _lst_text and "技能乙" not in _lst_text, _lst_text[:100])
    # 插件逐项开关：默认启用，可切换，禁用后 hook 不执行
    from sidecar.plugin_loader.loader import PluginLoader
    import sidecar.plugin_loader.loader as _ldm
    _orig_root = _ldm.PLUGINS_ROOT
    _ldm.PLUGINS_ROOT = TMP / "plugins"
    _ldm.PLUGINS_ROOT.mkdir(parents=True, exist_ok=True)
    try:
        pl = PluginLoader()
        pl._state_path = _ldm.PLUGINS_ROOT / "plugins_state.json"
        pd = _ldm.PLUGINS_ROOT / "demo"
        pd.mkdir()
        (pd / "manifest.json").write_text('{"name": "demo", "hooks": ["on_message"]}', encoding="utf-8")
        (pd / "plugin.py").write_text(
            "def on_message(ctx):\n    return 'CALLED'\n", encoding="utf-8")
        check("1g1 插件默认启用", pl.is_enabled("demo") is True)
        check("1g2 切换为禁用", pl.toggle_enabled("demo") is False)
        check("1g3 禁用状态生效", pl.is_enabled("demo") is False)
        check("1g4 列表附逐项状态", pl.list_installed()[0]["enabled"] is False)
        _r = asyncio.run(pl.execute_hook("on_message", {}))
        check("1g5 禁用插件 hook 不执行", _r is None, str(_r))
        pl.toggle_enabled("demo")
        _r2 = asyncio.run(pl.execute_hook("on_message", {}))
        check("1g6 重新启用后 hook 执行", _r2 is not None and _r2.get("result") == "CALLED", str(_r2))
        check("1g7 不存在的插件切换返回 None", pl.toggle_enabled("ghost") is None)
        # 卸载清理状态条目
        pl.uninstall("demo")
        check("1g8 卸载后状态条目清理", "demo" not in pl._read_state())
    finally:
        _ldm.PLUGINS_ROOT = _orig_root

    # ══ 2. 导出目录解析 ══
    from sidecar.exporter import resolve_export_dir
    wd = TMP / "user_work_dir"
    wd.mkdir()
    pid = store.create_project("proj", wd)

    d = resolve_export_dir(pid)
    # 项目工作目录经 resolve()（/var→/private/var 符号链接解析），对齐比较
    check("2a 导出目录空配置 → 项目工作目录", d == wd.resolve() or d == wd, str(d))

    cfgdir = TMP / "custom_exports"
    reload_config({"default_export_dir": str(cfgdir)})
    d2 = resolve_export_dir(pid)
    check("2b 配置非空 → 配置目录（自动创建）",
          (d2 == cfgdir or d2 == cfgdir.resolve()) and cfgdir.exists(), str(d2))

    reload_config({"default_export_dir": ""})

    # ══ 3. 会话导出 ══
    aid = store.add_agent_config(pid, "A", "main", model_name="qwen3.8")
    sid = store.create_session(pid, aid)
    store.save_message(pid, sid, aid, "user", "你好")
    store.save_message(pid, sid, aid, "assistant", "回复", tool_steps=[{"name": "list_dir", "ok": True, "summary": "3 项"}])

    from sidecar.exporter import export_session_md
    res = export_session_md(pid, sid)
    p = Path(res["path"])
    check("3a 会话导出文件存在", p.exists(), str(p))
    txt = p.read_text(encoding="utf-8")
    check("3b 导出含消息内容", "你好" in txt and "回复" in txt, txt[:200])
    check("3c 导出含工具步骤摘要", "list_dir" in txt, txt[:200])
    check("3d 导出在导出目录的 sessions/ 下", p.parent.name == "sessions", str(p.parent))

    try:
        export_session_md(pid, "no-such-session")
        check("3e 会话不存在 → ValueError", False, "未抛错")
    except ValueError:
        check("3e 会话不存在 → ValueError", True)

    # ══ 4. 交卷超长落盘 ══
    from sidecar.agent_engine.delegation import (
        parse_report, build_fallback_report, _finalize_summary, SUMMARY_MAX_LEN,
    )
    check("4a SUMMARY_MAX_LEN == 1000", SUMMARY_MAX_LEN == 1000, str(SUMMARY_MAX_LEN))

    tid = "task-m7"
    big_summary = "报" * 1200
    report = {"task_id": tid, "status": "success", "summary": big_summary, "artifacts": []}
    fin = _finalize_summary(pid, tid, report, "子 Agent 原始全文：" + big_summary)
    check("4b 超 1000 字 → summary 前 1000 字 + 路径标注",
          len(fin["summary"].split("\n[")[0]) == 1000 and "交卷全文已保存" in fin["summary"],
          fin["summary"][:80])
    saved = Path(fin.get("summary_saved_path", ""))
    check("4c 交卷全文落盘且含全文", saved.exists() and big_summary in saved.read_text(encoding="utf-8"),
          str(saved))
    check("4d 落盘目录在导出目录的 delegation_reports/ 下",
          saved.parent.name == "delegation_reports", str(saved.parent))

    short_report = {"task_id": tid, "status": "success", "summary": "短摘要", "artifacts": []}
    fin2 = _finalize_summary(pid, tid, short_report, "x")
    check("4e ≤1000 字原样回传不落盘",
          fin2["summary"] == "短摘要" and "summary_saved_path" not in fin2, str(fin2))

    # 兜底打包：超长全文保留（截断在落盘环节）
    fb = build_fallback_report("实" * 1300, tid)
    check("4f 兜底打包保留全文（>1000 不在此截断）",
          fb is not None and len(fb["summary"]) > 1000, str(len(fb["summary"]) if fb else None))

    # ══ 5. 附件解析器 ══
    from sidecar.attachments.parser import parse_attachment

    t, k = parse_attachment("笔记.txt", "会议纪要".encode("utf-8"))
    check("5a 纯文本解析", t == "会议纪要" and k == "text", f"{t}/{k}")
    t, k = parse_attachment("编码.md", "内容".encode("gbk"))
    check("5b GBK 文本可解析", t == "内容" and k == "text", f"{t}/{k}")
    t, k = parse_attachment("数据.csv", "a,b\n1,2\n3,4".encode())
    check("5c CSV 解析为表格文本", t is not None and "1 | 2" in t and k == "csv", f"{t}/{k}")
    t, k = parse_attachment("图.png", bytes([137, 80, 78, 71]))
    check("5d 图片 → 不识别仅标注", t is None and k == "image", f"{t}/{k}")
    t, k = parse_attachment("二进制.dat", bytes([0, 1, 255]))
    check("5e 无法解析 → binary", t is None and k == "binary", f"{t}/{k}")
    t, k = parse_attachment("损坏.pdf", b"not a pdf")
    check("5f 损坏 PDF → None 降级（不抛错）", t is None and k == "pdf", f"{t}/{k}")
    t, k = parse_attachment("损坏.docx", b"not a docx")
    check("5g 损坏 Word → None 降级", t is None and k == "docx", f"{t}/{k}")
    t, k = parse_attachment("损坏.xlsx", b"not a xlsx")
    check("5h 损坏 Excel → None 降级", t is None and k == "xlsx", f"{t}/{k}")

    # 构造真实最小 PDF / Word / Excel 验证解析器可用（依赖已装）
    try:
        from pypdf import PdfWriter
        import io as _io
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        buf = _io.BytesIO()
        w.write(buf)
        t, k = parse_attachment("空页.pdf", buf.getvalue())
        check("5i 真实 PDF 可解析（空页 → None/空文本降级）", k == "pdf", f"{t}/{k}")
    except ImportError:
        print("SKIP 5i pypdf 未安装")
    try:
        import docx as _docx
        doc = _docx.Document()
        doc.add_paragraph("表格数据ABC")
        buf = _io.BytesIO()
        doc.save(buf)
        t, k = parse_attachment("文档.docx", buf.getvalue())
        check("5j 真实 Word 解析出文本", t is not None and "表格数据ABC" in t and k == "docx", f"{t}/{k}")
    except ImportError:
        print("SKIP 5j python-docx 未安装")
    try:
        import openpyxl as _op
        wb = _op.Workbook()
        ws = wb.active
        ws.append(["名称", "数值"])
        ws.append(["苹果", 3])
        buf = _io.BytesIO()
        wb.save(buf)
        t, k = parse_attachment("表.xlsx", buf.getvalue())
        check("5k 真实 Excel 解析出文本", t is not None and "苹果" in t and k == "xlsx", f"{t}/{k}")
    except ImportError:
        print("SKIP 5k openpyxl 未安装")

    # ══ 6. 清理 ══
    import shutil
    shutil.rmtree(TMP, ignore_errors=True)
    check("6 临时目录已清理", not TMP.exists())

    print(f"\n===== TS-113 M7 专项: PASS={PASS} FAIL={FAIL} =====")
    if FAILURES:
        print("FAILED:", ", ".join(FAILURES))
        sys.exit(1)


if __name__ == "__main__":
    main()
