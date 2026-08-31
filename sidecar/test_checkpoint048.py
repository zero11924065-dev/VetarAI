"""checkpoint-048 专项单测：会话自动总结端点 + 聊天附件解析端点。
venv 内 python test_checkpoint048.py 直接跑（需 PYTHONPATH）。只输出 PASS/FAIL 摘要。
"""
import asyncio
import base64
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

PASS, FAIL = 0, 0
FAILURES = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"PASS  {name}")
    else:
        FAIL += 1
        FAILURES.append(name)
        print(f"FAIL  {name}  {detail}")


def main():
    TMP = Path(tempfile.mkdtemp(prefix="ck048_"))
    import sidecar.config as cfgmod
    cfgmod.get_config_path = lambda: TMP / "config.json"
    cfgmod._MEM = {}

    import sidecar.storage.store as store
    store.PROJECTS_ROOT = TMP / "projects"
    store._GDB = TMP / "projects" / "_global.db"
    store.PROJECTS_ROOT.mkdir(parents=True, exist_ok=True)

    from sidecar import app as appmod
    appmod.get_config = lambda: {"default_model": "qwen3.8", "network_switch": "auto"}

    from fastapi.testclient import TestClient
    client = TestClient(appmod.app)

    wd = TMP / "wd"
    wd.mkdir()
    pid = store.create_project("proj", wd)
    aid = store.add_agent_config(pid, "A", "main", model_name="qwen3.8")

    # ══ 1. 会话自动总结 ══
    # 1a 不存在的会话 → 404
    r = client.post(f"/api/sessions/no-such/summarize?project_id={pid}")
    check("1a 会话不存在 → 404", r.status_code == 404, str(r.status_code))

    # 1b 空会话（无消息）→ 404
    sid_empty = store.create_session(pid, aid)
    r = client.post(f"/api/sessions/{sid_empty}/summarize?project_id={pid}")
    check("1b 无消息 → 404", r.status_code == 404, str(r.status_code))

    # 1c 仅 system 消息（无可总结内容）→ 422
    sid_sys = store.create_session(pid, aid)
    store.save_message(pid, sid_sys, aid, "system", "系统消息")
    r = client.post(f"/api/sessions/{sid_sys}/summarize?project_id={pid}")
    check("1c 仅 system 消息 → 422", r.status_code == 422, str(r.status_code))

    # 1d 正常会话 → 打桩连接器生成总结 → 落盘 + 返回
    sid = store.create_session(pid, aid)
    store.save_message(pid, sid, aid, "user", "帮我查一下项目结构")
    store.save_message(pid, sid, aid, "assistant", "项目有 src 和 tests 两个目录，共 12 个文件。")

    class FakeConnector:
        def __init__(self):
            self.calls = 0
            self.last_msgs = None
        async def chat(self, model, messages, images=None):
            self.calls += 1
            self.last_msgs = messages
            return "总结：用户要求查看项目结构，助手确认有 src/tests 两目录共 12 文件。"

    fake = FakeConnector()
    appmod.get_ollama_connector = lambda: fake
    r = client.post(f"/api/sessions/{sid}/summarize?project_id={pid}",
                    json={"agent_id": aid, "model": "qwen3.8"})
    check("1d 正常总结 → 200", r.status_code == 200, r.text[:200])
    d = r.json()
    check("1e 返回总结与保存路径",
          d.get("ok") is True and "总结：" in d.get("summary", "") and d.get("saved_file"), str(d)[:200])
    check("1f 总结文件落盘", Path(d["saved_file"]).exists(), d.get("saved_file", ""))
    check("1g 连接器收到对话原文", fake.calls == 1 and fake.last_msgs
          and "帮我查一下项目结构" in fake.last_msgs[0]["content"], str(fake.last_msgs)[:150])

    # 1h 模型返回空 → 502
    class EmptyConnector:
        async def chat(self, model, messages, images=None):
            return "   "
    appmod.get_ollama_connector = lambda: EmptyConnector()
    r = client.post(f"/api/sessions/{sid}/summarize?project_id={pid}")
    check("1h 模型空返回 → 502", r.status_code == 502, str(r.status_code))

    # 1i 超长会话截断（8000 字上限）：构造超大会话，验证提示词含截断标注
    sid_big = store.create_session(pid, aid)
    store.save_message(pid, sid_big, aid, "user", "长" * 9000)

    class CaptureConnector:
        def __init__(self):
            self.prompt = ""
        async def chat(self, model, messages, images=None):
            self.prompt = messages[0]["content"]
            return "截断总结"
    cap = CaptureConnector()
    appmod.get_ollama_connector = lambda: cap
    r = client.post(f"/api/sessions/{sid_big}/summarize?project_id={pid}")
    check("1i 超大会话 → 提示词含截断标注",
          r.status_code == 200 and "已截断" in cap.prompt, cap.prompt[-100:])

    # ══ 2. 聊天附件解析端点 ══
    def b64(data: bytes) -> str:
        return base64.b64encode(data).decode("ascii")

    # 2a 文本文件
    r = client.post("/api/attachments/parse",
                    json={"name": "笔记.txt", "content_base64": b64("会议纪要ABC".encode("utf-8"))})
    check("2a 文本解析", r.status_code == 200 and r.json()["text"] == "会议纪要ABC"
          and r.json()["kind"] == "text", r.text[:120])

    # 2b CSV
    r = client.post("/api/attachments/parse",
                    json={"name": "数据.csv", "content_base64": b64("a,b\n1,2".encode())})
    check("2b CSV 解析", r.status_code == 200 and "1 | 2" in (r.json()["text"] or "")
          and r.json()["kind"] == "csv", r.text[:120])

    # 2c Word（真实构造）
    try:
        import docx as _docx
        import io as _io
        doc = _docx.Document()
        doc.add_paragraph("合同要点XYZ")
        buf = _io.BytesIO()
        doc.save(buf)
        r = client.post("/api/attachments/parse",
                        json={"name": "合同.docx", "content_base64": b64(buf.getvalue())})
        check("2c Word 解析", r.status_code == 200 and "合同要点XYZ" in (r.json()["text"] or "")
              and r.json()["kind"] == "docx", r.text[:120])
    except ImportError:
        print("SKIP 2c python-docx 未安装")

    # 2d Excel（真实构造）
    try:
        import openpyxl as _op
        import io as _io2
        wb = _op.Workbook()
        wb.active.append(["物料", "数量"])
        wb.active.append(["螺丝", 100])
        buf2 = _io2.BytesIO()
        wb.save(buf2)
        r = client.post("/api/attachments/parse",
                        json={"name": "清单.xlsx", "content_base64": b64(buf2.getvalue())})
        check("2d Excel 解析", r.status_code == 200 and "螺丝" in (r.json()["text"] or "")
              and r.json()["kind"] == "xlsx", r.text[:120])
    except ImportError:
        print("SKIP 2d openpyxl 未安装")

    # 2e 图片 → text=null kind=image（聊天图片走多模态，不走文本注入）
    r = client.post("/api/attachments/parse",
                    json={"name": "图.png", "content_base64": b64(bytes([137, 80, 78, 71]))})
    check("2e 图片 → text=null", r.status_code == 200 and r.json()["text"] is None
          and r.json()["kind"] == "image", r.text[:120])

    # 2f 无法解析的二进制 → text=null kind=binary
    r = client.post("/api/attachments/parse",
                    json={"name": "程序.bin", "content_base64": b64(bytes([0, 1, 2, 255]))})
    check("2f 二进制 → text=null binary", r.status_code == 200 and r.json()["text"] is None
          and r.json()["kind"] == "binary", r.text[:120])

    # 2g checkpoint-067 R-2 完整优先：聊天附件单文件上限放宽到 10MB，超 10MB 才 400
    #（旧 2MB 限制已放宽；圆桌仍保留 2MB 不受影响）
    r = client.post("/api/attachments/parse",
                    json={"name": "big.txt", "content_base64": b64(b"x" * (10 * 1024 * 1024 + 1))})
    check("2g 超 10MB → 400", r.status_code == 400, str(r.status_code))

    # 2h checkpoint-067 R-2 完整优先：单文件文本上限放宽到 20 万字符，
    # 4000 字远低于上限不截断（旧 3000 字截断已移除，律所分析材料不再被切断）
    r = client.post("/api/attachments/parse",
                    json={"name": "long.txt", "content_base64": b64(("字" * 4000).encode("utf-8"))})
    d = r.json()
    check("2h 4000 字不截断（完整优先）", r.status_code == 200 and len(d["text"]) == 4000
          and d.get("truncated") is False, str(d.get("truncated")))

    # 2i 编码非法 → 400
    r = client.post("/api/attachments/parse",
                    json={"name": "bad.txt", "content_base64": "!!!非法base64!!!"})
    check("2i 编码非法 → 400", r.status_code == 400, str(r.status_code))

    # ══ 3. 清理 ══
    import shutil
    shutil.rmtree(TMP, ignore_errors=True)
    check("3 临时目录已清理", not TMP.exists())

    print(f"\n===== checkpoint-048 专项: PASS={PASS} FAIL={FAIL} =====")
    if FAILURES:
        print("FAILED:", ", ".join(FAILURES))
        sys.exit(1)


if __name__ == "__main__":
    main()
